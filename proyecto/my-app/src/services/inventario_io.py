"""Exportar e importar el inventario en Excel (.xlsx) o CSV."""
import csv
import io
import unicodedata
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from services import productos

COLUMNAS = ["Código", "Nombre", "Categoría", "Precio", "Stock"]
CAMPOS = {"codigo", "nombre", "categoria", "precio", "stock"}

# Nombres de columna aceptados al importar (sin tildes ni mayúsculas)
_ALIAS = {
    "codigo": "codigo", "cod": "codigo", "sku": "codigo",
    "nombre": "nombre", "descripcion": "nombre", "producto": "nombre",
    "nombre / descripcion": "nombre", "nombre/descripcion": "nombre",
    "categoria": "categoria",
    "precio": "precio", "valor": "precio",
    "stock": "stock", "cantidad": "stock", "unidades": "stock",
}


# ---------- Utilidades ----------
def nombre_archivo(extension: str) -> str:
    return f"inventario_stockin_{datetime.now():%Y-%m-%d}.{extension}"


def _normalizar(texto) -> str:
    texto = str(texto or "").strip().lower()
    texto = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in texto if not unicodedata.combining(c))


def _texto(valor) -> str:
    """Convierte una celda a texto (Excel entrega 1990.0 en vez de 1990)."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


# ---------- Exportar ----------
def exportar_excel() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario"
    ws.append(COLUMNAS)

    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="E65100")
        celda.alignment = Alignment(horizontal="center")

    for p in productos.listar():
        ws.append([p["codigo"], p["nombre"], p["categoria"], p["precio"], p["stock"]])

    for (celda,) in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        celda.number_format = '"$"#,##0'

    for letra, ancho in zip("ABCDE", [14, 32, 20, 14, 10]):
        ws.column_dimensions[letra].width = ancho
    ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def exportar_csv() -> bytes:
    """CSV con ';' y BOM para que Excel en español lo abra con tildes y columnas bien."""
    buffer = io.StringIO()
    escritor = csv.writer(buffer, delimiter=";")
    escritor.writerow(COLUMNAS)
    for p in productos.listar():
        escritor.writerow([p["codigo"], p["nombre"], p["categoria"], p["precio"], p["stock"]])
    return buffer.getvalue().encode("utf-8-sig")


# ---------- Importar ----------
def _leer_filas(nombre: str, datos: bytes) -> list[list]:
    extension = nombre.lower().rsplit(".", 1)[-1]

    if extension == "xlsx":
        try:
            wb = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
        except Exception:
            raise ValueError("No se pudo abrir el Excel. ¿Está dañado?")
        return [list(fila) for fila in wb.active.iter_rows(values_only=True)]

    if extension == "csv":
        try:
            texto = datos.decode("utf-8-sig")
        except UnicodeDecodeError:
            texto = datos.decode("latin-1")
        primera = texto.splitlines()[0] if texto else ""
        separador = ";" if primera.count(";") >= primera.count(",") else ","
        return list(csv.reader(io.StringIO(texto), delimiter=separador))

    raise ValueError("Formato no soportado. Usa un archivo .xlsx o .csv.")


def importar(nombre: str, datos: bytes | None) -> dict:
    """Crea los productos nuevos y actualiza los que ya existen (mismo código).
    Devuelve {"creados", "actualizados", "errores"}."""
    if not datos:
        raise ValueError("No se pudo leer el archivo.")

    filas = [
        f for f in _leer_filas(nombre, datos)
        if f and any(_texto(c) for c in f)
    ]
    if len(filas) < 2:
        raise ValueError("El archivo no tiene productos.")

    encabezado = [_ALIAS.get(_normalizar(c)) for c in filas[0]]
    faltan = CAMPOS - set(encabezado)
    if faltan:
        raise ValueError(
            "Faltan columnas: " + ", ".join(sorted(faltan))
            + ". Usa como plantilla un archivo exportado desde StockIN."
        )

    existentes = {p["codigo"]: p for p in productos.listar()}
    vistos: set[str] = set()
    creados = actualizados = 0
    errores: list[str] = []

    for numero, fila in enumerate(filas[1:], start=2):
        valores = {
            campo: _texto(fila[i]) if i < len(fila) else ""
            for i, campo in enumerate(encabezado) if campo
        }
        codigo = valores["codigo"].upper()

        if codigo and codigo in vistos:
            errores.append(f"Fila {numero}: el código {codigo} está repetido en el archivo.")
            continue
        vistos.add(codigo)

        try:
            if codigo in existentes:
                productos.actualizar(
                    existentes[codigo]["id"], codigo, valores["nombre"],
                    valores["categoria"], valores["precio"], valores["stock"],
                )
                actualizados += 1
            else:
                productos.crear(
                    codigo, valores["nombre"], valores["categoria"],
                    valores["precio"], valores["stock"],
                )
                creados += 1
        except ValueError as err:
            errores.append(f"Fila {numero}: {err}")

    return {"creados": creados, "actualizados": actualizados, "errores": errores}